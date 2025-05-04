import os
import logging
from datetime import datetime
from fastapi import FastAPI, Request, HTTPException, UploadFile, File, Form, BackgroundTasks
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, RedirectResponse, HTMLResponse, Response, StreamingResponse
import uvicorn
from typing import List, Optional, Dict, Any
from app.models.database import create_tables
from app.services.llm_service import LLMService
from app.services.file_service import FileService
from app.utils.file_handler import FileHandler
from app.utils.text_processor import TextProcessor
from app.utils.validators import SplitSettings, QASettings, ExportSettings
from app.utils.batch_processor import BatchProcessor
from app.utils.quality_evaluator import QualityEvaluator
from dotenv import load_dotenv
import uuid
from fastapi import status
import sqlite3
import json
import chardet
from docx import Document
import PyPDF2
import re
from pydantic import BaseModel
import csv, io
from fastapi.exceptions import RequestValidationError
from app.i18n import LANGS

# 加载环境变量
load_dotenv()

# 配置日志
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('app.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# 创建FastAPI应用
app = FastAPI(
    title="Dataset-Bit",
    description="一个用于处理和生成高质量问答数据集的工具",
    version="1.0.0"
)

# 配置CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 挂载静态文件目录
app.mount("/static", StaticFiles(directory="frontend/static"), name="static")

# 配置模板
templates = Jinja2Templates(directory="frontend/templates")

# 初始化服务
llm_service = LLMService(api_key=os.getenv("OPENAI_API_KEY"))
file_service = FileService()
batch_processor = BatchProcessor(llm_service)
quality_evaluator = QualityEvaluator(llm_service)

# 创建数据库表
create_tables()

# 数据库连接
def get_db():
    conn = sqlite3.connect('dataset_bit.db')
    conn.row_factory = sqlite3.Row
    return conn

# 初始化数据库
def init_db():
    conn = get_db()
    c = conn.cursor()
    
    # 创建文件表
    c.execute('''
        CREATE TABLE IF NOT EXISTS files (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            filename TEXT NOT NULL,
            file_path TEXT NOT NULL,
            file_type TEXT NOT NULL,
            file_size INTEGER NOT NULL,
            upload_time TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            status TEXT DEFAULT 'pending'
        )
    ''')
    
    # 创建文本片段表
    c.execute('''
        CREATE TABLE IF NOT EXISTS segments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_id INTEGER,
            content TEXT NOT NULL,
            segment_index INTEGER NOT NULL,
            created_time TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (file_id) REFERENCES files (id)
        )
    ''')
    
    # 创建问答对表
    c.execute('''
        CREATE TABLE IF NOT EXISTS qa_pairs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            segment_id INTEGER,
            question TEXT NOT NULL,
            answer TEXT NOT NULL,
            quality_score REAL DEFAULT 0,
            created_time TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (segment_id) REFERENCES segments (id)
        )
    ''')
    
    conn.commit()
    conn.close()

# 初始化数据库
init_db()

def init_settings_table():
    conn = get_db()
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS settings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            api_base TEXT,
            api_key TEXT,
            model_name TEXT,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            language TEXT DEFAULT 'zh',
            theme TEXT DEFAULT 'light'
        )
    ''')
    c.execute('SELECT COUNT(*) FROM settings')
    if c.fetchone()[0] == 0:
        c.execute('INSERT INTO settings (api_base, api_key, model_name, language, theme) VALUES ("", "", "", "zh", "light")')
    conn.commit()
    conn.close()

@app.get("/", include_in_schema=False)
async def root():
    return RedirectResponse(url="/file.html")

@app.get("/")
async def index(request: Request):
    """首页/仪表盘"""
    # 获取统计数据
    total_datasets = 4  # TODO: 从数据库获取实际数量
    total_qa_pairs = 42  # TODO: 从数据库获取实际数量
    total_exports = 5  # TODO: 从数据库获取实际数量
    
    # 获取最近活动
    activities = [
        {
            "description": "使用AI模型 'qwen-plus' 生成了18个问答对",
            "time": "2025-04-18 14:50"
        },
        {
            "description": "导入文档 数据资产管理实践白皮书 (5.0版) .pdf 到数据集 Data asset dataset",
            "time": "2025-04-18 14:48"
        },
        {
            "description": "导入文档 数据资产管理实践白皮书 (5.0版) .pdf 到数据集 Data asset dataset",
            "time": "2025-04-18 14:48"
        },
        {
            "description": "创建数据集: Data asset dataset",
            "time": "2025-04-18 14:47"
        },
        {
            "description": "使用AI模型 'qwen-plus' 生成了24个问答对",
            "time": "2025-04-18 14:44"
        }
    ]
    
    return templates.TemplateResponse(
        "index.html",
        {
            "request": request,
            "total_datasets": total_datasets,
            "total_qa_pairs": total_qa_pairs,
            "total_exports": total_exports,
            "activities": activities
        }
    )

@app.get("/upload")
async def upload_page(request: Request):
    """上传文档页面"""
    return templates.TemplateResponse("upload.html", {"request": request})

@app.get("/qa")
async def qa_page(request: Request):
    """问答生成页面"""
    return templates.TemplateResponse("qa.html", {"request": request})

@app.get("/datasets")
async def datasets_page(request: Request):
    """数据集管理页面"""
    return templates.TemplateResponse("datasets.html", {"request": request})

@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    """上传文件"""
    try:
        logger.info(f"开始处理文件上传: {file.filename}")
        
        # 检查文件类型
        allowed_types = ['txt', 'md', 'docx', 'pdf']
        file_ext = file.filename.split('.')[-1].lower()
        if file_ext not in allowed_types:
            logger.warning(f"不支持的文件类型: {file_ext}")
            raise HTTPException(status_code=400, detail="不支持的文件类型")
        
        # 保存文件
        file_path = os.path.join("uploads", file.filename)
        with open(file_path, "wb") as f:
            content = await file.read()
            f.write(content)
        
        # 记录到数据库
        conn = get_db()
        c = conn.cursor()
        file_size = os.path.getsize(file_path)
        c.execute(
            "INSERT INTO files (filename, file_path, file_type, file_size) VALUES (?, ?, ?, ?)",
            (file.filename, file_path, file_ext, file_size)
        )
        file_id = c.lastrowid
        conn.commit()
        conn.close()
        
        logger.info("文件上传处理完成")
        return {"status": "success", "file_id": file_id, "filename": file.filename}
    except Exception as e:
        logger.error(f"文件上传失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/split")
async def split_text(file_id: int, method: str = "paragraph", min_length: int = 100, max_length: int = 2000):
    """分割文本"""
    try:
        conn = get_db()
        c = conn.cursor()
        
        # 获取文件信息
        c.execute("SELECT * FROM files WHERE id = ?", (file_id,))
        file = c.fetchone()
        if not file:
            raise HTTPException(status_code=404, detail="文件不存在")
        
        # 读取文件内容
        with open(file['file_path'], 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 根据方法分割文本
        segments = []
        if method == "paragraph":
            segments = content.split('\n\n')
        elif method == "heading":
            # 简单的标题分割逻辑
            segments = [s.strip() for s in content.split('\n#') if s.strip()]
        else:
            raise HTTPException(status_code=400, detail="不支持的分割方法")
        
        # 过滤和保存分割结果
        for i, segment in enumerate(segments):
            if min_length <= len(segment) <= max_length:
                c.execute(
                    "INSERT INTO segments (file_id, content, segment_index) VALUES (?, ?, ?)",
                    (file_id, segment, i)
                )
        
        conn.commit()
        conn.close()
        return {"status": "success", "segments_count": len(segments)}
    except Exception as e:
        logger.error(f"文本分割失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/generate-questions")
async def generate_questions(settings: QASettings):
    """生成问题"""
    try:
        questions = await llm_service.generate_questions(settings.text, settings)
        logger.info(f"问题生成成功: {len(questions)}个问题")
        
        return {
            "status": "success",
            "message": "问题生成成功",
            "data": {
                "questions": questions,
                "count": len(questions)
            }
        }
    except Exception as e:
        logger.error(f"问题生成失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/generate-answers")
async def generate_answers(settings: QASettings):
    """生成答案"""
    try:
        answers = []
        for qa_pair in settings.qa_pairs:
            answer = await llm_service.generate_answers(
                qa_pair["question"],
                qa_pair["context"],
                settings
            )
            answers.append({
                "question": qa_pair["question"],
                "answer": answer
            })
        
        logger.info(f"答案生成成功: {len(answers)}个答案")
        
        return {
            "status": "success",
            "message": "答案生成成功",
            "data": {
                "qa_pairs": answers,
                "count": len(answers)
            }
        }
    except Exception as e:
        logger.error(f"答案生成失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/evaluate-quality")
async def evaluate_quality(qa_pairs: List[Dict[str, str]], context: str = None):
    """评估数据集质量"""
    try:
        # 评估数据集质量
        metrics = await quality_evaluator.evaluate_dataset(qa_pairs, context)
        
        # 生成质量报告
        report = quality_evaluator.get_quality_report(metrics)
        
        return {
            "status": "success",
            "message": "质量评估完成",
            "data": {
                "metrics": metrics,
                "report": report
            }
        }
    except Exception as e:
        logger.error(f"质量评估失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/export")
async def export_dataset(format: str = "alpaca", include_metadata: bool = True):
    """导出数据集"""
    try:
        conn = get_db()
        c = conn.cursor()
        
        # 获取所有问答对
        c.execute("""
            SELECT qa.*, s.content as context, f.filename
            FROM qa_pairs qa
            JOIN segments s ON qa.segment_id = s.id
            JOIN files f ON s.file_id = f.id
            ORDER BY qa.created_time DESC
        """)
        qa_pairs = [dict(row) for row in c.fetchall()]
        
        # 根据格式导出
        if format == "alpaca":
            export_data = []
            for qa in qa_pairs:
                item = {
                    "instruction": qa['question'],
                    "input": qa['context'],
                    "output": qa['answer']
                }
                if include_metadata:
                    item["metadata"] = {
                        "source_file": qa['filename'],
                        "quality_score": qa['quality_score'],
                        "created_time": qa['created_time']
                    }
                export_data.append(item)
            
            # 保存到文件
            export_path = os.path.join("exports", f"dataset_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json")
            with open(export_path, 'w', encoding='utf-8') as f:
                json.dump(export_data, f, ensure_ascii=False, indent=2)
            
            return {"status": "success", "export_path": export_path}
        else:
            raise HTTPException(status_code=400, detail="不支持的导出格式")
    except Exception as e:
        logger.error(f"数据集导出失败: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/datasets")
async def get_datasets():
    """获取所有文件清单及问答对数量"""
    try:
        conn = get_db()
        c = conn.cursor()
        c.execute("""
            SELECT f.id, f.filename, COUNT(q.id) as qa_count, f.created_at as created_at
            FROM files f
            LEFT JOIN qa_pairs q ON f.id = q.file_id
            GROUP BY f.id, f.filename, f.created_at
            ORDER BY f.created_at DESC
        """)
        files = [dict(row) for row in c.fetchall()]
        conn.close()
        return files
    except Exception as e:
        logger.error(f"获取文件清单失败: {str(e)}")
        raise HTTPException(status_code=500, detail="获取文件清单失败")

@app.get("/api/datasets/{dataset_id}")
async def get_dataset(dataset_id: int):
    """获取数据集详情"""
    try:
        # 获取数据集基本信息
        dataset = await db.fetch_one("""
            SELECT 
                d.id,
                d.name,
                d.created_at,
                COUNT(qa.id) as qa_count,
                COALESCE(AVG(qa.quality_score), 0) as quality_score
            FROM datasets d
            LEFT JOIN qa_pairs qa ON d.id = qa.dataset_id
            WHERE d.id = :dataset_id
            GROUP BY d.id, d.name, d.created_at
        """, {"dataset_id": dataset_id})
        
        if not dataset:
            raise HTTPException(status_code=404, detail="数据集不存在")
        
        # 获取质量评估指标
        metrics = await db.fetch_one("""
            SELECT 
                COALESCE(AVG(accuracy), 0) as accuracy,
                COALESCE(AVG(completeness), 0) as completeness,
                COALESCE(AVG(relevance), 0) as relevance,
                COALESCE(AVG(clarity), 0) as clarity,
                COALESCE(AVG(quality_score), 0) as total_score
            FROM qa_pairs
            WHERE dataset_id = :dataset_id
        """, {"dataset_id": dataset_id})
        
        # 获取问答对预览
        qa_pairs = await db.fetch_all("""
            SELECT question, answer
            FROM qa_pairs
            WHERE dataset_id = :dataset_id
            ORDER BY id DESC
            LIMIT 3
        """, {"dataset_id": dataset_id})
        
        return {
            **dataset,
            "metrics": metrics,
            "qa_pairs": qa_pairs
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"获取数据集详情失败: {str(e)}")
        raise HTTPException(status_code=500, detail="获取数据集详情失败")

@app.post("/api/datasets/{dataset_id}/evaluate")
async def evaluate_dataset(dataset_id: int):
    """评估数据集质量"""
    try:
        # 获取数据集的所有问答对
        qa_pairs = await db.fetch_all("""
            SELECT id, question, answer
            FROM qa_pairs
            WHERE dataset_id = :dataset_id
        """, {"dataset_id": dataset_id})
        
        if not qa_pairs:
            raise HTTPException(status_code=404, detail="数据集不存在或为空")
        
        # 批量评估问答对质量
        evaluator = QualityEvaluator()
        for qa in qa_pairs:
            quality_score = await evaluator.evaluate_qa_pair(qa["question"], qa["answer"])
            
            # 更新问答对质量分数
            await db.execute("""
                UPDATE qa_pairs
                SET quality_score = :quality_score
                WHERE id = :qa_id
            """, {
                "qa_id": qa["id"],
                "quality_score": quality_score
            })
        
        return {"message": "数据集评估完成"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"评估数据集失败: {str(e)}")
        raise HTTPException(status_code=500, detail="评估数据集失败")

@app.get("/api/datasets/{dataset_id}/export")
def export_dataset(dataset_id: int, format: str = 'alpaca', type: str = 'json'):
    conn = get_db()
    c = conn.cursor()
    # 查询qa_pairs表，假设有file_id字段关联
    c.execute("SELECT question, answer FROM qa_pairs WHERE file_id=?", (dataset_id,))
    rows = c.fetchall()
    conn.close()
    # 格式化数据
    data = []
    if format == 'alpaca':
        for q, a in rows:
            data.append({"instruction": q, "input": "", "output": a})
    elif format == 'sharegpt':
        for q, a in rows:
            data.append({"conversations": [{"from": "human", "value": q}, {"from": "gpt", "value": a}]})
    else:
        data = [{"question": q, "answer": a} for q, a in rows]
    # 导出类型
    if type == 'json':
        import json
        content = json.dumps(data, ensure_ascii=False, indent=2)
        return StreamingResponse(io.BytesIO(content.encode('utf-8')), media_type='application/json', headers={"Content-Disposition": f"attachment; filename=dataset_{dataset_id}_{format}.json"})
    elif type == 'csv':
        output = io.StringIO()
        writer = csv.writer(output)
        if format == 'alpaca':
            writer.writerow(['instruction', 'input', 'output'])
            for item in data:
                writer.writerow([item['instruction'], item['input'], item['output']])
        elif format == 'sharegpt':
            writer.writerow(['question', 'answer'])
            for item in data:
                q = item['conversations'][0]['value']
                a = item['conversations'][1]['value']
                writer.writerow([q, a])
        else:
            writer.writerow(['question', 'answer'])
            for item in data:
                writer.writerow([item['question'], item['answer']])
        return StreamingResponse(io.BytesIO(output.getvalue().encode('utf-8')), media_type='text/csv', headers={"Content-Disposition": f"attachment; filename=dataset_{dataset_id}_{format}.csv"})
    elif type == 'md':
        md = ''
        if format == 'alpaca':
            for item in data:
                md += f"### 指令\n{item['instruction']}\n\n### 输出\n{item['output']}\n\n---\n"
        elif format == 'sharegpt':
            for item in data:
                q = item['conversations'][0]['value']
                a = item['conversations'][1]['value']
                md += f"**Q:** {q}\n\n**A:** {a}\n\n---\n"
        else:
            for item in data:
                md += f"Q: {item['question']}\nA: {item['answer']}\n\n---\n"
        return StreamingResponse(io.BytesIO(md.encode('utf-8')), media_type='text/markdown', headers={"Content-Disposition": f"attachment; filename=dataset_{dataset_id}_{format}.md"})
    else:
        return {"status": "error", "message": "不支持的导出类型"}

@app.delete("/api/datasets/{dataset_id}")
async def delete_dataset(dataset_id: int):
    """删除数据集"""
    try:
        # 检查数据集是否存在
        dataset = await db.fetch_one("""
            SELECT id
            FROM datasets
            WHERE id = :dataset_id
        """, {"dataset_id": dataset_id})
        
        if not dataset:
            raise HTTPException(status_code=404, detail="数据集不存在")
        
        # 删除数据集及其关联的问答对
        await db.execute("""
            DELETE FROM qa_pairs
            WHERE dataset_id = :dataset_id
        """, {"dataset_id": dataset_id})
        
        await db.execute("""
            DELETE FROM datasets
            WHERE id = :dataset_id
        """, {"dataset_id": dataset_id})
        
        return {"message": "数据集已删除"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"删除数据集失败: {str(e)}")
        raise HTTPException(status_code=500, detail="删除数据集失败")

@app.post("/api/batch/process")
async def process_batch(
    text_segments: List[str],
    settings: QASettings,
    background_tasks: BackgroundTasks
):
    """开始批量处理"""
    batch_id = str(uuid.uuid4())
    
    # 在后台处理
    background_tasks.add_task(
        batch_processor.process_batch,
        text_segments=text_segments,
        settings=settings,
        batch_id=batch_id
    )
    
    return {
        "status": "started",
        "batch_id": batch_id,
        "message": "批量处理已开始"
    }

@app.get("/api/batch/progress/{batch_id}")
async def get_batch_progress(batch_id: str):
    """获取处理进度"""
    progress = await batch_processor.get_progress(batch_id)
    if not progress:
        raise HTTPException(status_code=404, detail="找不到处理进度")
    return progress

@app.post("/api/batch/resume/{batch_id}")
async def resume_batch(
    batch_id: str,
    text_segments: List[str],
    settings: QASettings,
    background_tasks: BackgroundTasks
):
    """继续处理"""
    # 在后台继续处理
    background_tasks.add_task(
        batch_processor.resume_processing,
        batch_id=batch_id,
        text_segments=text_segments,
        settings=settings
    )
    
    return {
        "status": "resumed",
        "batch_id": batch_id,
        "message": "批量处理已继续"
    }

@app.delete("/api/batch/{batch_id}")
async def cleanup_batch(batch_id: str):
    """清理处理数据"""
    await batch_processor.cleanup(batch_id)
    return {"status": "success", "message": "清理完成"}

@app.get("/files")
async def files_page(request: Request):
    """文件列表页面"""
    # 获取上传目录中的所有文件
    upload_dir = os.path.join("app", "static", "uploads")
    files = []
    
    if os.path.exists(upload_dir):
        for filename in os.listdir(upload_dir):
            file_path = os.path.join(upload_dir, filename)
            if os.path.isfile(file_path):
                # 获取文件信息
                stat = os.stat(file_path)
                # 从文件名中提取原始文件名（去掉时间戳前缀）
                original_filename = '_'.join(filename.split('_')[2:])
                
                files.append({
                    "id": filename,  # 使用文件名作为ID
                    "filename": original_filename,
                    "size": format_file_size(stat.st_size),
                    "upload_time": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
                    "status": "待处理",  # TODO: 从数据库获取实际状态
                    "status_color": "warning"  # TODO: 根据状态设置颜色
                })
    
    # 按上传时间倒序排序
    files.sort(key=lambda x: x["upload_time"], reverse=True)
    
    # 添加状态的多语言映射
    status_map = {
        '待处理': t['status_pending'],
        '已分块': t['status_chunked'],
        '已完成': t['status_done'],
        'pending': t['status_pending'],
        'chunked': t['status_chunked'],
        'done': t['status_done'],
    }
    for f in files:
        f['status_translated'] = status_map.get(f['status'], f['status'])
    
    return templates.TemplateResponse(
        "files.html",
        {
            "request": request,
            "files": files
        }
    )

def format_file_size(size):
    """格式化文件大小"""
    for unit in ['B', 'KB', 'MB', 'GB']:
        if size < 1024.0:
            return f"{size:.1f} {unit}"
        size /= 1024.0
    return f"{size:.1f} TB"

@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    """全局异常处理"""
    logger.error(f"发生错误: {str(exc)}", exc_info=True)
    return JSONResponse(
        status_code=500,
        content={
            "status": "error",
            "message": "服务器内部错误",
            "detail": str(exc)
        }
    )

@app.exception_handler(RequestValidationError)
async def validation_exception_handler(request: Request, exc: RequestValidationError):
    logger.error(f"422参数校验失败: url={request.url} detail={exc.errors()} body={await request.body()}")
    return JSONResponse(
        status_code=422,
        content={"detail": exc.errors(), "body": exc.body if hasattr(exc, 'body') else None}
    )

# 新增：多页面路由
@app.get("/file.html", response_class=HTMLResponse)
async def file_page(request: Request):
    """文件管理页面"""
    lang = request.query_params.get('lang') or request.cookies.get('lang') or 'zh'
    t = LANGS.get(lang, LANGS['zh'])
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM files ORDER BY created_at DESC")
    files = [dict(row) for row in cursor.fetchall()]
    conn.close()
    status_map = {
        '待处理': t['status_pending'],
        '已分块': t['status_chunked'],
        '已完成': t['status_done'],
        'pending': t['status_pending'],
        'chunked': t['status_chunked'],
        'done': t['status_done'],
    }
    for f in files:
        f['status_translated'] = status_map.get(f['status'], f['status'])
    return templates.TemplateResponse("file.html", {
        "request": request,
        "files": files,
        "page": "file",
        "t": t,
        "lang": lang
    })

@app.get("/chunk.html", response_class=HTMLResponse)
async def chunk_page(request: Request):
    """分块管理页面"""
    lang = request.query_params.get('lang') or request.cookies.get('lang') or 'zh'
    t = LANGS[lang]
    return templates.TemplateResponse("chunk.html", {
        "request": request,
        "page": "chunk",
        "lang": lang,
        "t": t
    })

@app.get("/dataset.html", response_class=HTMLResponse)
async def dataset_page(request: Request):
    """数据导出页面"""
    lang = request.query_params.get('lang') or request.cookies.get('lang') or 'zh'
    t = LANGS[lang]
    return templates.TemplateResponse("dataset.html", {
        "request": request,
        "page": "dataset",
        "lang": lang,
        "t": t
    })

@app.get("/settings.html", response_class=HTMLResponse)
async def settings_page(request: Request):
    """系统设置页面"""
    lang = request.query_params.get('lang') or request.cookies.get('lang') or 'zh'
    t = LANGS[lang]
    return templates.TemplateResponse("settings.html", {
        "request": request,
        "page": "settings",
        "lang": lang,
        "t": t
    })

split_progress = {}

class SplitParams(BaseModel):
    method: str = "paragraph"
    block_size: int = 1000
    overlap: int = 15

# 在分块接口中限制block_size范围
    block_size = max(100, min(block_size, 5000))

@app.post("/api/files/{file_id}/split")
async def split_file(file_id: int, params: SplitParams, background_tasks: BackgroundTasks):
    method = params.method
    block_size = params.block_size
    overlap = params.overlap
    split_progress[file_id] = {"current": 0, "total": 1, "status": "processing"}
    def do_split():
        import sqlite3
        conn = sqlite3.connect('dataset_bit.db')
        c = conn.cursor()
        c.execute("SELECT file_path, file_type FROM files WHERE id=?", (file_id,))
        row = c.fetchone()
        if not row:
            split_progress[file_id] = {"current": 0, "total": 1, "status": "error"}
            return
        file_path, file_type = row
        content = read_file_content(file_path, file_type)
        blocks = split_content(content, method, block_size, overlap)
        total = len(blocks)
        split_progress[file_id] = {"current": 0, "total": total, "status": "processing"}
        c.execute("DELETE FROM text_segments WHERE file_id=?", (file_id,))
        for i, block in enumerate(blocks):
            c.execute("INSERT INTO text_segments (file_id, content, segment_index) VALUES (?, ?, ?)", (file_id, block, i))
            split_progress[file_id] = {"current": i+1, "total": total, "status": "processing"}
        conn.commit()
        c.execute("UPDATE files SET status='已分块' WHERE id=?", (file_id,))
        conn.commit()
        conn.close()
        split_progress[file_id] = {"current": total, "total": total, "status": "done"}
    background_tasks.add_task(do_split)
    return {"status": "started"}

@app.get("/api/files/{file_id}/split_progress")
async def split_progress_api(file_id: int):
    return split_progress.get(file_id, {"current": 0, "total": 1, "status": "not_started"})

def read_file_content(file_path, file_type):
    if file_type in ['txt', 'md']:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            with open(file_path, 'rb') as f:
                raw = f.read()
                result = chardet.detect(raw)
                encoding = result['encoding'] or 'utf-8'
                return raw.decode(encoding, errors='ignore')
    elif file_type == 'docx':
        doc = Document(file_path)
        return '\n'.join([p.text for p in doc.paragraphs])
    elif file_type == 'pdf':
        reader = PyPDF2.PdfReader(file_path)
        return '\n'.join([page.extract_text() or '' for page in reader.pages])
    else:
        return ''

def split_by_length_within_block(block, block_size, overlap_rate):
    result = []
    overlap_rate = max(0, min(overlap_rate, 99))
    step = int(block_size * (1 - overlap_rate / 100))
    if step < 1:
        step = 1
    i = 0
    while i < len(block):
        sub_block = block[i:i+block_size]
        if sub_block.strip():
            result.append(sub_block)
        i += step
    return result

def split_by_heading(content):
    # 以标题为分块起点，返回每个标题块
    content = content.replace('\r\n', '\n').replace('\r', '\n')
    pattern = r'(^#+\s.*$)'
    lines = content.split('\n')
    blocks = []
    current_block = []
    for line in lines:
        if re.match(pattern, line.strip()):
            if current_block:
                blocks.append('\n'.join(current_block))
                current_block = []
        current_block.append(line)
    if current_block:
        blocks.append('\n'.join(current_block))
    return [b.strip() for b in blocks if b.strip()]

def split_content(content, method, block_size, overlap):
    content = content.replace('\r\n', '\n').replace('\r', '\n')
    paragraph_splitter = r'(?:\n\s*){2,}'
    if method == "paragraph":
        blocks = [seg for seg in re.split(paragraph_splitter, content) if seg.strip()]
    elif method == "heading":
        # 每个标题块合并为一个分块
        blocks = split_by_heading(content)
    elif method == "table":
        blocks = re.split(r'\n\|', content)
        blocks = [b.strip() for b in blocks if b.strip()]
    elif method == "auto":
        # 智能递归分层：先按标题分块，再对每个标题块按段落分块
        heading_blocks = split_by_heading(content)
        blocks = []
        for hblock in heading_blocks:
            blocks.extend([seg for seg in re.split(paragraph_splitter, hblock) if seg.strip()])
    else:
        blocks = [seg for seg in re.split(paragraph_splitter, content) if seg.strip()]
    # 对每个块，超长才切分，绝不跨块拼接
    final_blocks = []
    for block in blocks:
        if len(block) <= block_size:
            final_blocks.append(block)
        else:
            final_blocks.extend(split_by_length_within_block(block, block_size, overlap))
    return final_blocks

@app.get("/api/files")
async def get_files():
    """获取所有文件列表"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, filename as file_name FROM files ORDER BY id DESC")
        files = [dict(row) for row in cursor.fetchall()]
        conn.close()
        return {"files": files}
    except Exception as e:
        logger.error(f"获取文件列表失败: {str(e)}")
        raise HTTPException(status_code=500, detail="获取文件列表失败")

@app.get("/api/files/{file_id}/chunks")
async def get_file_chunks(file_id: int, page: int = 1, page_size: int = 10):
    """获取文件的分块列表"""
    try:
        conn = get_db()
        cursor = conn.cursor()
        
        # 获取总记录数
        cursor.execute("SELECT COUNT(*) FROM text_segments WHERE file_id = ?", (file_id,))
        total = cursor.fetchone()[0]
        
        # 获取分页数据
        offset = (page - 1) * page_size
        cursor.execute("""
            SELECT id, segment_index as chunk_index, content
            FROM text_segments
            WHERE file_id = ?
            ORDER BY segment_index
            LIMIT ? OFFSET ?
        """, (file_id, page_size, offset))
        
        chunks = []
        for row in cursor.fetchall():
            chunk = dict(row)
            content = chunk.get('content', '')
            chunk['full_content'] = content
            if len(content) > 95:
                chunk['content'] = content[:95] + '...'
            chunks.append(chunk)
        conn.close()
        
        return {
            "chunks": chunks,
            "total": total,
            "page": page,
            "page_size": page_size
        }
    except Exception as e:
        logger.error(f"获取分块列表失败: {str(e)}")
        raise HTTPException(status_code=500, detail="获取分块列表失败")

@app.delete("/api/chunks/{chunk_id}")
async def delete_chunk(chunk_id: int):
    try:
        logger.info(f"收到删除分块请求 chunk_id={chunk_id}")
        conn = get_db()
        cursor = conn.cursor()
        # 先查找分块对应的 segment_id
        cursor.execute("SELECT id FROM text_segments WHERE id = ?", (chunk_id,))
        row = cursor.fetchone()
        if not row:
            conn.close()
            return {"status": "error", "message": "分块不存在或已删除"}
        segment_id = row[0]
        # 删除 qa_pairs 表中关联的问答对
        cursor.execute("DELETE FROM qa_pairs WHERE segment_id = ?", (segment_id,))
        # 删除 text_segments 表中的分块
        cursor.execute("DELETE FROM text_segments WHERE id = ?", (chunk_id,))
        affected = cursor.rowcount
        conn.commit()
        conn.close()
        logger.info(f"删除分块结果: affected={affected}")
        if affected == 0:
            return {"status": "error", "message": "分块不存在或已删除"}
        return {"status": "success", "message": "分块及相关问答对已删除"}
    except Exception as e:
        logger.error(f"删除分块失败: {str(e)}")
        raise HTTPException(status_code=500, detail="删除分块失败")

@app.get("/api/settings")
async def get_settings():
    """获取系统设置"""
    try:
        conn = get_db()
        c = conn.cursor()
        c.execute("SELECT api_base, api_key, model_name, language, theme FROM settings LIMIT 1")
        row = c.fetchone()
        conn.close()
        if row:
            return {
                "status": "success",
                "data": {
                    "api_base": row[0] or "",
                    "api_key": row[1] or "",
                    "model_name": row[2] or "",
                    "language": row[3] or "zh",
                    "theme": row[4] or "light"
                }
            }
        return {
            "status": "success",
            "data": {
                "api_base": "",
                "api_key": "",
                "model_name": "",
                "language": "zh",
                "theme": "light"
            }
        }
    except Exception as e:
        logger.error(f"获取设置失败: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/settings")
async def save_settings(request: Request):
    """保存系统设置"""
    try:
        data = await request.json()
        api_base = data.get('api_base', '')
        api_key = data.get('api_key', '')
        model_name = data.get('model_name', '')
        language = data.get('language', 'zh')
        theme = data.get('theme', 'light')
        conn = get_db()
        c = conn.cursor()
        c.execute("UPDATE settings SET api_base=?, api_key=?, model_name=?, language=?, theme=?, updated_at=CURRENT_TIMESTAMP WHERE id=1", (api_base, api_key, model_name, language, theme))
        conn.commit()
        conn.close()
        return JSONResponse({"status": "success"})
    except Exception as e:
        return JSONResponse({"status": "error", "message": str(e)}, status_code=500)

@app.post("/api/settings/test")
async def test_settings(data: dict):
    """测试API连接"""
    try:
        import openai
        openai.api_key = data.get("api_key", "")
        base_url = data.get("api_base", "")
        if base_url:
            if not base_url.endswith("/"):
                base_url += "/"
            openai.base_url = base_url
        
        # 测试连接
        response = openai.chat.completions.create(
            model=data.get("model_name", "gpt-3.5-turbo"),
            messages=[{"role": "user", "content": "Hello"}],
            max_tokens=5
        )
        
        return {"status": "success", "message": "连接成功"}
    except Exception as e:
        logger.error(f"测试连接失败: {str(e)}")
        return {"status": "error", "message": str(e)}

@app.post("/api/generate-qa")
async def generate_qa(data: dict):
    try:
        segments = data.get("segments", [])
        num_pairs = int(data.get("num_pairs", 3))
        file_id = int(data.get("file_id", 0))
        lang = data.get("lang", "zh")
        if not segments or not num_pairs or not file_id:
            return {"status": "error", "message": "参数不完整"}
        # 获取大模型设置
        conn = get_db()
        c = conn.cursor()
        c.execute("SELECT api_base, api_key, model_name FROM settings LIMIT 1")
        row = c.fetchone()
        conn.close()
        if not row:
            return {"status": "error", "message": "未配置大模型参数"}
        api_base, api_key, model_name = row
        import openai
        openai.api_key = api_key
        base_url = api_base
        if not base_url.endswith("/"):
            base_url += "/"
        openai.base_url = base_url
        import json
        total_qa = 0
        conn = get_db()
        c = conn.cursor()
        for seg in segments:
            if isinstance(seg, dict):
                segment_id = seg.get('id', 0)
                seg_content = seg.get('content', '')
            else:
                segment_id = 0
                seg_content = seg
            if lang == 'en':
                prompt = f"""Based on the following text, generate {num_pairs} QA pairs. Each pair should include a question and an answer.\nEnsure the questions are diverse, including both open-ended and factual ones. Answers should be accurate, complete, and based on the text.\nReturn the result in JSON format as follows:\n[{{\"question\": \"Question 1\", \"answer\": \"Answer 1\"}}, {{\"question\": \"Question 2\", \"answer\": \"Answer 2\"}}]\n\nText:\n{seg_content}\n\nPlease reply in English."""
            else:
                prompt = f"""基于以下文本内容，生成{num_pairs}个问答对。每个问答对应包含问题和答案。\n请确保问题多样化，包括开放性问题和事实性问题。答案应该准确、完整且基于文本内容。\n请以JSON格式返回结果，格式为：\n[{{\"question\": \"问题1\", \"answer\": \"答案1\"}}, {{\"question\": \"问题2\", \"answer\": \"答案2\"}}]\n\n文本内容：\n{seg_content}"""
            completion = openai.chat.completions.create(
                model=model_name,
                messages=[{"role": "user", "content": prompt}],
                timeout=30
            )
            content = completion.choices[0].message.content.strip()
            try:
                qa_list = json.loads(content)
            except Exception:
                import re
                match = re.search(r'\[.*\]', content, re.DOTALL)
                if match:
                    qa_list = json.loads(match.group(0))
                else:
                    continue
            for qa in qa_list:
                c.execute("INSERT INTO qa_pairs (segment_id, question, answer, quality_scores, file_id) VALUES (?, ?, ?, ?, ?)", (segment_id, qa.get("question", ""), qa.get("answer", ""), '{}', file_id))
                total_qa += 1
        conn.commit()
        conn.close()
        return {"status": "success", "count": total_qa}
    except Exception as e:
        return {"status": "error", "message": str(e)}

# 禁用所有缓存
@app.middleware("http")
async def no_cache_middleware(request, call_next):
    response = await call_next(request)
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

@app.get("/api/datasets_export")
def export_datasets(ids: Optional[str] = None, format: str = 'alpaca', type: str = 'json'):
    if not ids:
        return {"status": "error", "message": "缺少导出文件ID"}
    try:
        file_ids = [int(id) for id in ids.split(',')]
        conn = get_db()
        c = conn.cursor()
        # 查询所有选中文件的问答对
        placeholders = ','.join('?' * len(file_ids))
        c.execute(f"""
            SELECT q.question, q.answer, f.filename 
            FROM qa_pairs q
            JOIN files f ON q.file_id = f.id
            WHERE q.file_id IN ({placeholders})
            ORDER BY f.id, q.id
        """, file_ids)
        rows = c.fetchall()
        conn.close()
        # 格式化数据
        data = []
        if format == 'alpaca':
            for q, a, filename in rows:
                data.append({
                    "instruction": q,
                    "input": "",
                    "output": a
                })
        elif format == 'sharegpt':
            for q, a, filename in rows:
                data.append({
                    "conversations": [
                        {"from": "human", "value": q},
                        {"from": "gpt", "value": a}
                    ]
                })
        else:
            data = [{"question": q, "answer": a, "source_file": filename} for q, a, filename in rows]
        # 导出类型
        if type == 'json':
            import json
            content = json.dumps(data, ensure_ascii=False, indent=2)
            return StreamingResponse(
                io.BytesIO(content.encode('utf-8')),
                media_type='application/json',
                headers={"Content-Disposition": f"attachment; filename=dataset_export_{format}.json"}
            )
        elif type == 'csv':
            output = io.StringIO()
            writer = csv.writer(output)
            if format == 'alpaca':
                writer.writerow(['instruction', 'input', 'output'])
                for item in data:
                    writer.writerow([
                        item['instruction'],
                        item['input'],
                        item['output']
                    ])
            elif format == 'sharegpt':
                writer.writerow(['question', 'answer'])
                for item in data:
                    q = item['conversations'][0]['value']
                    a = item['conversations'][1]['value']
                    writer.writerow([q, a])
            else:
                writer.writerow(['question', 'answer', 'source_file'])
                for item in data:
                    writer.writerow([item['question'], item['answer'], item['source_file']])
            return StreamingResponse(
                io.BytesIO(output.getvalue().encode('utf-8')),
                media_type='text/csv',
                headers={"Content-Disposition": f"attachment; filename=dataset_export_{format}.csv"}
            )
        elif type == 'md':
            md = ''
            if format == 'alpaca':
                for item in data:
                    md += f"### 指令\n{item['instruction']}\n\n### 输出\n{item['output']}\n\n---\n"
            elif format == 'sharegpt':
                for item in data:
                    q = item['conversations'][0]['value']
                    a = item['conversations'][1]['value']
                    md += f"**Q:** {q}\n\n**A:** {a}\n\n---\n"
            else:
                for item in data:
                    md += f"Q: {item['question']}\nA: {item['answer']}\n\n来源文件: {item['source_file']}\n\n---\n"
            return StreamingResponse(
                io.BytesIO(md.encode('utf-8')),
                media_type='text/markdown',
                headers={"Content-Disposition": f"attachment; filename=dataset_export_{format}.md"}
            )
        else:
            return {"status": "error", "message": "不支持的导出类型"}
    except Exception as e:
        logger.error(f"导出数据集失败: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

def migrate_add_file_id_to_qa_pairs():
    conn = get_db()
    c = conn.cursor()
    # 检查file_id字段是否已存在
    c.execute("PRAGMA table_info(qa_pairs)")
    columns = [row[1] for row in c.fetchall()]
    if 'file_id' not in columns:
        c.execute("ALTER TABLE qa_pairs ADD COLUMN file_id INTEGER")
        conn.commit()
    # 补全历史数据：通过segment_id查segments表，写入file_id
    c.execute("SELECT id, segment_id FROM qa_pairs WHERE file_id IS NULL OR file_id=''")
    qa_rows = c.fetchall()
    for qa_id, segment_id in qa_rows:
        if segment_id:
            c.execute("SELECT file_id FROM segments WHERE id=?", (segment_id,))
            row = c.fetchone()
            if row and row[0]:
                c.execute("UPDATE qa_pairs SET file_id=? WHERE id=?", (row[0], qa_id))
    conn.commit()
    conn.close()

# 启动时自动迁移
migrate_add_file_id_to_qa_pairs()

@app.post("/api/files/{file_id}/delete")
async def delete_file(file_id: int):
    try:
        conn = get_db()
        c = conn.cursor()
        # 删除 qa_pairs 表中相关数据
        c.execute("DELETE FROM qa_pairs WHERE file_id=?", (file_id,))
        # 删除 text_segments 表中相关数据
        c.execute("DELETE FROM text_segments WHERE file_id=?", (file_id,))
        # 删除 files 表中相关数据
        c.execute("DELETE FROM files WHERE id=?", (file_id,))
        conn.commit()
        conn.close()
        return {"status": "success", "message": "文件及相关数据已删除"}
    except Exception as e:
        logger.error(f"删除文件失败: {str(e)}")
        return {"status": "error", "message": str(e)}

@app.get("/api/chunks/{segment_id}/qa")
async def get_chunk_qa(segment_id: int):
    try:
        conn = get_db()
        c = conn.cursor()
        c.execute("SELECT id, question, answer FROM qa_pairs WHERE segment_id=? ORDER BY id ASC", (segment_id,))
        qa_list = [{"id": row[0], "question": row[1], "answer": row[2]} for row in c.fetchall()]
        conn.close()
        return {"status": "success", "data": qa_list, "count": len(qa_list)}
    except Exception as e:
        logger.error(f"获取分块问答对失败: {str(e)}")
        return {"status": "error", "message": str(e)}

@app.post("/api/qa/{qa_id}/update")
async def update_qa(qa_id: int, data: dict):
    try:
        question = data.get("question", "").strip()
        answer = data.get("answer", "").strip()
        if not question or not answer:
            return {"status": "error", "message": "问题和答案不能为空"}
        conn = get_db()
        c = conn.cursor()
        c.execute("UPDATE qa_pairs SET question=?, answer=? WHERE id=?", (question, answer, qa_id))
        conn.commit()
        conn.close()
        return {"status": "success"}
    except Exception as e:
        logger.error(f"更新问答对失败: {str(e)}")
        return {"status": "error", "message": str(e)}

@app.post("/api/qa/{qa_id}/delete")
async def delete_qa(qa_id: int):
    try:
        conn = get_db()
        c = conn.cursor()
        c.execute("DELETE FROM qa_pairs WHERE id=?", (qa_id,))
        conn.commit()
        conn.close()
        return {"status": "success"}
    except Exception as e:
        logger.error(f"删除问答对失败: {str(e)}")
        return {"status": "error", "message": str(e)}

if __name__ == "__main__":
    uvicorn.run(
        "main:app",
        host=os.getenv("APP_HOST", "0.0.0.0"),
        port=int(os.getenv("APP_PORT", 8000)),
        reload=True
    ) 